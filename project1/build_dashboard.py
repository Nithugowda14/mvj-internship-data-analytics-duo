"""
Online Retail II -> interactive Excel dashboard builder.

Usage:
    python build_dashboard.py path/to/online_retail_II.xlsx  output.xlsx
Optional env var RAW_CACHE=raw.pkl re-uses a pickled copy of the raw data (the
1M-row Excel read takes ~90 s).  Requires: pandas, numpy, openpyxl.
"""
import sys, os, pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DFont
from openpyxl.drawing.line import LineProperties

SRC = sys.argv[1] if len(sys.argv) > 1 else '/mnt/user-data/uploads/online_retail_II.xlsx'
OUT = sys.argv[2] if len(sys.argv) > 2 else '/mnt/user-data/outputs/Online_Retail_Dashboard.xlsx'
TOP_PRODUCTS = 500
SNAPSHOT = pd.Timestamp('2011-12-10')

# ----------------------------------------------------------------------------------
# 1. LOAD + CLEAN + AGGREGATE
# ----------------------------------------------------------------------------------
def load_raw():
    cache = os.environ.get('RAW_CACHE')
    if cache and os.path.exists(cache):
        return pd.read_pickle(cache)
    a = pd.read_excel(SRC, sheet_name='Year 2009-2010'); a['src'] = '0910'
    b = pd.read_excel(SRC, sheet_name='Year 2010-2011'); b['src'] = '1011'
    return pd.concat([a, b], ignore_index=True)

df = load_raw()
log = {'raw_rows': len(df)}
ov = (df.src == '1011') & (df.InvoiceDate < '2010-12-10')       # 1-9 Dec 2010 appear in both sheets
log['overlap'] = int(ov.sum()); df = df[~ov].copy()
df['Invoice'] = df.Invoice.astype(str)
df['StockCode'] = df.StockCode.astype(str).str.strip()
df['Country'] = df.Country.astype(str).str.strip()
NONMERCH = {'POST','DOT','M','D','S','B','C2','BANK CHARGES','ADJUST','ADJUST2','AMAZONFEE','CRUK','TEST001','TEST002','PADS'}
sc = df.StockCode
is_non = sc.str.upper().isin(NONMERCH) | sc.str.lower().str.startswith('gift_') | df.Invoice.str.startswith('A')
log['nonmerch'] = int(is_non.sum()); df = df[~is_non]
m = df.Price <= 0;                                    log['price0'] = int(m.sum()); df = df[~m]
df['cancel'] = df.Invoice.str.startswith('C')
m = (~df.cancel) & (df.Quantity <= 0);                log['negqty'] = int(m.sum()); df = df[~m]
m = df.cancel & (df.Quantity > 0);                    log['poscancel'] = int(m.sum()); df = df[~m]
log['clean'] = len(df)
df['Rev'] = df.Quantity * df.Price
df['Year'] = df.InvoiceDate.dt.year; df['MonthNum'] = df.InvoiceDate.dt.month
df['Month'] = df.InvoiceDate.dt.to_period('M').dt.to_timestamp()
df['Weekday'] = df.InvoiceDate.dt.dayofweek + 1; df['Hour'] = df.InvoiceDate.dt.hour
sales, ret = df[~df.cancel], df[df.cancel]
log['no_cust_rows'] = int(df['Customer ID'].isna().sum())
log['no_cust_rev_share'] = float(df[df['Customer ID'].isna()].Rev.sum() / df.Rev.sum())
log['dmin'], log['dmax'] = df.InvoiceDate.min(), df.InvoiceDate.max()

g = ['Month', 'Year', 'MonthNum', 'Country']
fm = sales.groupby(g).agg(Sales=('Rev','sum'), Orders=('Invoice','nunique'), Units=('Quantity','sum')).reset_index()
fr = ret.groupby(g).agg(Returns=('Rev', lambda s: -s.sum())).reset_index()
fm = fm.merge(fr, on=g, how='outer').fillna(0)[['Month','Year','MonthNum','Country','Sales','Returns','Orders','Units']].sort_values(['Month','Country'])

prod_rev = df.groupby('StockCode').Rev.sum().sort_values(ascending=False)
top = prod_rev.head(TOP_PRODUCTS).index
desc = (df[df.StockCode.isin(top) & df.Description.notna()].groupby('StockCode').Description
        .agg(lambda s: s.str.strip().value_counts().index[0]))
pm = pd.DataFrame({'StockCode': top}); pm['ProdID'] = np.arange(1, len(pm) + 1)
pm['Description'] = pm.StockCode.map(desc).fillna('(no description)')
log['prod_cov'] = float(prod_rev.head(TOP_PRODUCTS).sum() / prod_rev.sum())
fp = (df[df.StockCode.isin(top)].merge(pm[['StockCode','ProdID']], on='StockCode')
        .groupby(['ProdID','Country','Year']).agg(Revenue=('Rev','sum'), Units=('Quantity','sum')).reset_index())
ft = sales.groupby(['Year','Country','Weekday','Hour']).Rev.sum().reset_index().rename(columns={'Rev':'Sales'})

c = df[df['Customer ID'].notna()].copy(); c['cid'] = c['Customer ID'].astype(int)
cs = c[~c.cancel]
cust = cs.groupby('cid').agg(Orders=('Invoice','nunique'), First=('InvoiceDate','min'), Last=('InvoiceDate','max')).reset_index()
yr = c.pivot_table(index='cid', columns='Year', values='Rev', aggfunc='sum', fill_value=0).reset_index()
yr.columns = ['cid'] + [f'Rev{y}' for y in yr.columns[1:]]
for y in (2009, 2010, 2011):
    if f'Rev{y}' not in yr: yr[f'Rev{y}'] = 0.0
cty = (c.groupby(['cid','Country']).Rev.sum().reset_index().sort_values(['cid','Rev'], ascending=[True, False])
         .drop_duplicates('cid')[['cid','Country']])
cust = cust.merge(yr, on='cid').merge(cty, on='cid')
cust['First'] = cust.First.dt.normalize(); cust['Last'] = cust.Last.dt.normalize()
cust = cust.sort_values('cid').reset_index(drop=True)

country_order = (fm.assign(n=fm.Sales - fm.Returns).groupby('Country').n.sum().sort_values(ascending=False).index.tolist())
print('Data ready:', log)

# ----------------------------------------------------------------------------------
# 2. STYLES
# ----------------------------------------------------------------------------------
NAVY, TEAL, AMBER, GREY, LIGHT, WHITE = '1F2A44', '0E9AA7', 'F4A261', '8A94A6', 'EEF1F6', 'FFFFFF'
CORAL, PURPLE, GREEN = 'E76F51', '6C5CE7', '2A9D8F'
FN = 'Arial'
def font(sz=10, b=False, color='000000', i=False): return Font(name=FN, size=sz, bold=b, color=color, italic=i)
def fill(c): return PatternFill('solid', start_color=c, end_color=c)
thin = Side(style='thin', color='D5DAE3'); med = Side(style='medium', color=TEAL)
INPUT_BLUE = '0000FF'

wb = Workbook()
wsD = wb.active; wsD.title = 'Dashboard'
wsN = wb.create_sheet('Notes'); wsC = wb.create_sheet('Calc'); wsU = wb.create_sheet('Customers')
wsFM = wb.create_sheet('Fact_Monthly'); wsFP = wb.create_sheet('Fact_Product'); wsFT = wb.create_sheet('Fact_Time')
wsL = wb.create_sheet('Lists')
for w, col in ((wsD, NAVY), (wsN, TEAL), (wsC, GREY), (wsU, GREY), (wsFM, GREY), (wsFP, GREY), (wsFT, GREY), (wsL, GREY)):
    w.sheet_properties.tabColor = col

def header(ws, row, labels, col0=1):
    for j, t in enumerate(labels):
        x = ws.cell(row=row, column=col0 + j, value=t)
        x.font = font(10, True, WHITE); x.fill = fill(NAVY); x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# ---- Lists
wsL['A1'] = 'Countries'; wsL['C1'] = 'Years'; wsL['E1'] = 'Metrics'; wsL['G1'] = 'Weekdays'; wsL['I1'] = 'Segments'
for ref in ('A1','C1','E1','G1','I1'): wsL[ref].font = font(10, True)
countries = ['All Countries'] + country_order
for i, v in enumerate(countries): wsL.cell(row=2 + i, column=1, value=v)
for i, v in enumerate(['All Years', 2009, 2010, 2011]): wsL.cell(row=2 + i, column=3, value=v)
METRICS = ['Net Revenue', 'Orders', 'Units Sold', 'Avg Order Value']
for i, v in enumerate(METRICS): wsL.cell(row=2 + i, column=5, value=v)
WD = ['Mon','Tue','Wed','Thu','Fri','Sat','Sun']
for i, v in enumerate(WD): wsL.cell(row=2 + i, column=7, value=v)
SEGS = ['Champions', 'Loyal', 'Recent / New', 'At Risk', 'Hibernating / Lost']
for i, v in enumerate(SEGS): wsL.cell(row=2 + i, column=9, value=v)
for col in 'ACEGI': wsL.column_dimensions[col].width = 22
for row in wsL.iter_rows(min_row=2):
    for x in row: x.font = font(10)
NCTRY = len(countries)

# ---- Fact tables
def dump(ws, frame, fmts=None):
    header(ws, 1, list(frame.columns))
    for r in frame.itertuples(index=False):
        ws.append([x.to_pydatetime() if isinstance(x, pd.Timestamp) else (float(x) if isinstance(x, (np.floating,)) else (int(x) if isinstance(x, (np.integer,)) else x)) for x in r])
    for row in ws.iter_rows(min_row=2):
        for x in row: x.font = font(10)
    for k, v in (fmts or {}).items():
        for x in ws[k][1:]: x.number_format = v
    ws.freeze_panes = 'A2'
    for j in range(1, len(frame.columns) + 1): ws.column_dimensions[L(j)].width = 15
dump(wsFM, fm, {'A': 'mmm-yy', 'E': '#,##0.00', 'F': '#,##0.00', 'G': '#,##0', 'H': '#,##0'}); nFM = len(fm) + 1
dump(wsFP, fp, {'D': '#,##0.00', 'E': '#,##0'}); nFP = len(fp) + 1
dump(wsFT, ft, {'E': '#,##0.00'}); nFT = len(ft) + 1
# product master on Fact_Product G:I
header(wsFP, 1, ['ProdID', 'StockCode', 'Description'], 7)
for i, r in enumerate(pm.itertuples(index=False)):
    wsFP.cell(row=2 + i, column=7, value=int(r.ProdID)); wsFP.cell(row=2 + i, column=8, value=str(r.StockCode)); wsFP.cell(row=2 + i, column=9, value=str(r.Description))
    for j in (7, 8, 9): wsFP.cell(row=2 + i, column=j).font = font(10)
wsFP.column_dimensions['I'].width = 42; wsFP.column_dimensions['H'].width = 14

FMr = lambda col: f"Fact_Monthly!${col}$2:${col}${nFM}"
FPr = lambda col: f"Fact_Product!${col}$2:${col}${nFP}"
FTr = lambda col: f"Fact_Time!${col}$2:${col}${nFT}"

# ---- Customers (data + RFM formulas)
nU = len(cust) + 1
uh = ['CustomerID','Country','First Order','Last Order','Orders','Net Rev 2009','Net Rev 2010','Net Rev 2011','Total Net Rev',
      'Recency (days)','R Score','F Score','Segment','Rev (selected year)','Rank key']
header(wsU, 1, uh)
for i, r in enumerate(cust.itertuples(index=False)):
    rw = 2 + i
    wsU.append([int(r.cid), r.Country, r.First.to_pydatetime(), r.Last.to_pydatetime(), int(r.Orders), float(r.Rev2009), float(r.Rev2010), float(r.Rev2011),
        f"=SUM(F{rw}:H{rw})",
        f"=Calc!$C$5-D{rw}",
        f"=IF(J{rw}<=Calc!$C$11,4,IF(J{rw}<=Calc!$C$12,3,IF(J{rw}<=Calc!$C$13,2,1)))",
        f"=1+(E{rw}>=Calc!$C$16)+(E{rw}>=Calc!$C$17)+(E{rw}>=Calc!$C$18)",
        f'=IF(AND(K{rw}=4,L{rw}>=3),"Champions",IF(AND(K{rw}=3,L{rw}>=3),"Loyal",IF(K{rw}>=3,"Recent / New",IF(L{rw}>=3,"At Risk","Hibernating / Lost"))))',
        f"=IF(AND(Calc!$C$2<=2009,Calc!$C$3>=2009),F{rw},0)+IF(AND(Calc!$C$2<=2010,Calc!$C$3>=2010),G{rw},0)+IF(AND(Calc!$C$2<=2011,Calc!$C$3>=2011),H{rw},0)",
        f"=IF(AND(N{rw}>0,OR(Calc!$C$4=\"*\",B{rw}=Calc!$C$4)),N{rw}+ROW()*0.000000001,0)"])
f10 = font(10)
for row in wsU.iter_rows(min_row=2):
    for x in row: x.font = f10
for k, v in {'C': 'dd-mmm-yy', 'D': 'dd-mmm-yy', 'F': '#,##0.00', 'G': '#,##0.00', 'H': '#,##0.00', 'I': '#,##0.00', 'N': '#,##0.00', 'O': '0.000000'}.items():
    for x in wsU[k][1:]: x.number_format = v
wsU.freeze_panes = 'A2'; wsU.auto_filter.ref = f"A1:O{nU}"
for j in range(1, 16): wsU.column_dimensions[L(j)].width = 15
wsU.column_dimensions['M'].width = 20
Ur = lambda col: f"Customers!${col}$2:${col}${nU}"

# ----------------------------------------------------------------------------------
# 3. CALC SHEET (all dashboard logic lives here as live formulas)
# ----------------------------------------------------------------------------------
C = wsC
C.column_dimensions['A'].width = 3; C.column_dimensions['B'].width = 40
for col in 'CDEFGHIJ': C.column_dimensions[col].width = 16
def sec(row, text):
    x = C.cell(row=row, column=2, value=text); x.font = font(11, True, WHITE); x.fill = fill(TEAL)
    for j in range(3, 11): C.cell(row=row, column=j).fill = fill(TEAL)
def lab(row, text): C.cell(row=row, column=2, value=text).font = font(10)
def val(row, formula, fmt=None, col=3, inp=False):
    x = C.cell(row=row, column=col, value=formula); x.font = font(10, color=INPUT_BLUE if inp else '000000')
    if fmt: x.number_format = fmt
    return x
sec(1, 'x'); C['B1'] = 'CALC ENGINE - all dashboard numbers are live formulas on the Fact_* sheets. Blue cells = editable assumptions.'; C['B1'].font = font(11, True, WHITE)
lab(2, 'Year filter - low');  val(2, '=IFERROR(VALUE(Dashboard!$H$4),0)')
lab(3, 'Year filter - high'); val(3, '=IFERROR(VALUE(Dashboard!$H$4),9999)')
lab(4, 'Country criterion ("*" = all)'); val(4, '=IF(Dashboard!$C$4="All Countries","*",Dashboard!$C$4)')
lab(5, 'Snapshot date for recency (day after last order)'); val(5, SNAPSHOT.to_pydatetime(), 'dd-mmm-yyyy', inp=True)
lab(6, 'Last COMPLETE month of 2011 (for YoY comparison)'); val(6, 11, '0', inp=True)
lab(7, 'Year label'); val(7, '=IFERROR(TEXT(VALUE(Dashboard!$H$4),"0"),"All years")')
lab(8, 'Country label'); val(8, '=IF(Dashboard!$C$4="All Countries","All countries",Dashboard!$C$4)')
sec(10, 'RFM recency cut-offs (days since last order, quartiles of all customers)')
lab(11, 'Score 4 if recency <= (25th pct)'); val(11, f'=PERCENTILE({Ur("J")},0.25)', '0')
lab(12, 'Score 3 if recency <= (50th pct)'); val(12, f'=PERCENTILE({Ur("J")},0.5)', '0')
lab(13, 'Score 2 if recency <= (75th pct)'); val(13, f'=PERCENTILE({Ur("J")},0.75)', '0')
sec(15, 'RFM frequency cut-offs (number of orders) - editable')
lab(16, 'F score +1 at >= orders'); val(16, 2, '0', inp=True)
lab(17, 'F score +1 again at >= orders'); val(17, 4, '0', inp=True)
lab(18, 'F score +1 again at >= orders'); val(18, 8, '0', inp=True)

crit = f'{FMr("D")},$C$4,{FMr("B")},">="&$C$2,{FMr("B")},"<="&$C$3'
sec(20, 'KPI calculations (follow Country + Year filters)')
lab(21, 'Gross sales (£)');       val(21, f'=SUMIFS({FMr("E")},{crit})', '#,##0.00')
lab(22, 'Returns / cancellations (£)'); val(22, f'=SUMIFS({FMr("F")},{crit})', '#,##0.00')
lab(23, 'Net revenue (£)');       val(23, '=C21-C22', '#,##0.00')
lab(24, 'Orders (unique invoices)'); val(24, f'=SUMIFS({FMr("G")},{crit})', '#,##0')
lab(25, 'Units sold');            val(25, f'=SUMIFS({FMr("H")},{crit})', '#,##0')
lab(26, 'Avg order value (gross sales / orders)'); val(26, '=IF(C24=0,0,C21/C24)', '#,##0.00')
lab(27, 'Return rate (returns / gross sales)'); val(27, '=IF(C21=0,0,C22/C21)', '0.0%')
lab(28, 'Active customers (with Customer ID)'); val(28, f'=COUNTIF({Ur("O")},">0")', '#,##0')
lab(29, 'Active customers - all countries'); val(29, f'=COUNTIF({Ur("N")},">0")', '#,##0')
lab(30, 'Net revenue from identified customers (£)'); val(30, f'=SUMIFS({Ur("N")},{Ur("B")},$C$4,{Ur("N")},">0")', '#,##0.00')
lab(31, 'Revenue per active customer (£)'); val(31, '=IF(C28=0,0,C30/C28)', '#,##0')
lab(32, 'Items per order'); val(32, '=IF(C24=0,0,C25/C24)', '0.0')

def ycrit(yr, extra=''): return f'{FMr("D")},$C$4,{FMr("B")},{yr},{FMr("C")},"<="&$C$6'
sec(34, 'Year-over-year comparison, 2011 vs 2010 (same months Jan..last complete month)')
for r_, t_, col_, yr_ in ((35,'Net revenue 2011 (£)','E',2011),(36,'Net revenue 2010 (£)','E',2010)):
    lab(r_, t_); val(r_, f'=SUMIFS({FMr("E")},{ycrit(yr_)})-SUMIFS({FMr("F")},{ycrit(yr_)})', '#,##0')
lab(37, 'Orders 2011'); val(37, f'=SUMIFS({FMr("G")},{ycrit(2011)})', '#,##0')
lab(38, 'Orders 2010'); val(38, f'=SUMIFS({FMr("G")},{ycrit(2010)})', '#,##0')
lab(39, 'Units 2011');  val(39, f'=SUMIFS({FMr("H")},{ycrit(2011)})', '#,##0')
lab(40, 'Units 2010');  val(40, f'=SUMIFS({FMr("H")},{ycrit(2010)})', '#,##0')
lab(41, 'Gross sales 2011 (£)'); val(41, f'=SUMIFS({FMr("E")},{ycrit(2011)})', '#,##0')
lab(42, 'Gross sales 2010 (£)'); val(42, f'=SUMIFS({FMr("E")},{ycrit(2010)})', '#,##0')
lab(43, 'Net revenue YoY');  val(43, '=IF(C36=0,"n/a",C35/C36-1)', '0.0%')
lab(44, 'Orders YoY');       val(44, '=IF(C38=0,"n/a",C37/C38-1)', '0.0%')
lab(45, 'Units YoY');        val(45, '=IF(C40=0,"n/a",C39/C40-1)', '0.0%')
lab(46, 'Avg order value YoY'); val(46, '=IF(OR(C37=0,C38=0,C42=0),"n/a",(C41/C37)/(C42/C38)-1)', '0.0%')
def sub(d):
    return (f'=IF(AND($C$2=2011,$C$3=2011),TEXT({d},"+0.0%;-0.0%")&" YoY (Jan-Nov)",IF(AND($C$2=2010,$C$3=2010),"Full year 2010",'
            f'IF(AND($C$2=2009,$C$3=2009),"Dec 2009 only (partial)","Dec 2009 - Dec 2011")))')
sec(47, 'Card sub-texts')
lab(48, 'Net revenue sub'); val(48, sub('C43'))
lab(49, 'Orders sub');      val(49, sub('C44'))
lab(50, 'AOV sub');         val(50, sub('C46'))
lab(51, 'Units sub');       val(51, sub('C45'))
lab(52, 'Customers sub');   val(52, '=IF($C$4="*","Unique buyers with a Customer ID",TEXT(C28/MAX(C29,1),"0.0%")&" of all customers")')
lab(53, 'Return sub');      val(53, '="Returns value: £"&TEXT(C22,"#,##0")')
lab(54, 'Rev/customer sub'); val(54, 'Identified customers only')
lab(55, 'Items/order sub');  val(55, 'Units sold / orders')

# Monthly table
R_M = 58; sec(R_M - 1, 'Monthly series (follows Country filter; ignores Year filter)')
for j, t in enumerate(['Month','Gross sales','Returns','Net revenue','Orders','Units','Avg order value', None, 'Label']):
    if t: C.cell(row=R_M, column=2 + j, value=t).font = font(10, True)
C.cell(row=R_M, column=10, value='=Dashboard!$L$4').font = font(10, True)   # selected-metric header (series name)
C.cell(row=R_M, column=10).alignment = Alignment(horizontal='left')
# columns: B month, C sales, D returns, E net, F orders, G units, H aov, I label, J selected metric
C.cell(row=R_M, column=9, value='Label').font = font(10, True)
for i in range(25):
    r_ = R_M + 1 + i
    C.cell(row=r_, column=2, value=(pd.Timestamp('2009-12-01').to_pydatetime() if i == 0 else f'=EDATE(B{r_-1},1)')).number_format = 'mmm-yy'
    C.cell(row=r_, column=3, value=f'=SUMIFS({FMr("E")},{FMr("A")},$B{r_},{FMr("D")},$C$4)').number_format = '#,##0'
    C.cell(row=r_, column=4, value=f'=SUMIFS({FMr("F")},{FMr("A")},$B{r_},{FMr("D")},$C$4)').number_format = '#,##0'
    C.cell(row=r_, column=5, value=f'=C{r_}-D{r_}').number_format = '#,##0'
    C.cell(row=r_, column=6, value=f'=SUMIFS({FMr("G")},{FMr("A")},$B{r_},{FMr("D")},$C$4)').number_format = '#,##0'
    C.cell(row=r_, column=7, value=f'=SUMIFS({FMr("H")},{FMr("A")},$B{r_},{FMr("D")},$C$4)').number_format = '#,##0'
    C.cell(row=r_, column=8, value=f'=IF(F{r_}=0,0,C{r_}/F{r_})').number_format = '#,##0.00'
    C.cell(row=r_, column=9, value=f'=TEXT(B{r_},"mmm yy")')
    C.cell(row=r_, column=10, value=f'=CHOOSE(MATCH(Dashboard!$L$4,Lists!$E$2:$E$5,0),E{r_},F{r_},G{r_},H{r_})').number_format = '#,##0'
R_M1, R_M2 = R_M + 1, R_M + 25

# YoY table
R_Y = 88; sec(R_Y - 1, 'Year-over-year net revenue by calendar month (follows Country filter)')
for j, t in enumerate(['Month #', 'Month']): C.cell(row=R_Y, column=2 + j, value=t).font = font(10, True)
C.cell(row=R_Y, column=4, value=2010).font = font(10, True); C.cell(row=R_Y, column=5, value=2011).font = font(10, True)
MN = ['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
for i in range(12):
    r_ = R_Y + 1 + i
    C.cell(row=r_, column=2, value=i + 1); C.cell(row=r_, column=3, value=MN[i])
    for col in (4, 5):
        cl = L(col)
        C.cell(row=r_, column=col, value=f'=SUMIFS({FMr("E")},{FMr("B")},{cl}${R_Y},{FMr("C")},$B{r_},{FMr("D")},$C$4)-SUMIFS({FMr("F")},{FMr("B")},{cl}${R_Y},{FMr("C")},$B{r_},{FMr("D")},$C$4)').number_format = '#,##0'
R_Y1, R_Y2 = R_Y + 1, R_Y + 12

# Country ranking
R_K = 104; sec(R_K - 1, 'Country ranking (follows Year filter; ignores Country filter; UK excluded from ranking)')
for j, t in enumerate(['Country', 'Net revenue', 'Rank key']): C.cell(row=R_K, column=2 + j, value=t).font = font(10, True)
nk = NCTRY - 1
for i in range(nk):
    r_ = R_K + 1 + i
    C.cell(row=r_, column=2, value=f'=Lists!A{3+i}')
    yc = f'{FMr("D")},$B{r_},{FMr("B")},">="&$C$2,{FMr("B")},"<="&$C$3'
    C.cell(row=r_, column=3, value=f'=SUMIFS({FMr("E")},{yc})-SUMIFS({FMr("F")},{yc})').number_format = '#,##0'
    C.cell(row=r_, column=4, value=f'=IF(B{r_}="United Kingdom",-1000000000,C{r_})+ROW()*0.000001').number_format = '0.000000'
R_K1, R_K2 = R_K + 1, R_K + nk
lab(R_K2 + 1, 'UK share of net revenue'); val(R_K2 + 1, f'=IFERROR(INDEX(C{R_K1}:C{R_K2},MATCH("United Kingdom",B{R_K1}:B{R_K2},0))/SUM(C{R_K1}:C{R_K2}),0)', '0.0%')
UKSHARE = f'Calc!$C${R_K2+1}'
R_KT = R_K2 + 4
for j, t in enumerate(['Rank', 'Index', 'Country', 'Net revenue']): C.cell(row=R_KT, column=2 + j, value=t).font = font(10, True)
for k in range(10):
    r_ = R_KT + 1 + k
    C.cell(row=r_, column=2, value=k + 1)
    C.cell(row=r_, column=3, value=f'=MATCH(LARGE($D${R_K1}:$D${R_K2},B{r_}),$D${R_K1}:$D${R_K2},0)')
    C.cell(row=r_, column=4, value=f'=IF(INDEX($C${R_K1}:$C${R_K2},C{r_})<=0,"-",INDEX($B${R_K1}:$B${R_K2},C{r_}))')
    C.cell(row=r_, column=5, value=f'=MAX(INDEX($C${R_K1}:$C${R_K2},C{r_}),0)').number_format = '£#,##0,"k"'
R_KT1, R_KT2 = R_KT + 1, R_KT + 10

# Product ranking
R_P = R_KT2 + 4; sec(R_P - 1, 'Product ranking (top-500 product universe; follows Country + Year filters)')
for j, t in enumerate(['ProdID', 'StockCode', 'Description', 'Net revenue', 'Rank key', 'Units']): C.cell(row=R_P, column=2 + j, value=t).font = font(10, True)
for i in range(TOP_PRODUCTS):
    r_ = R_P + 1 + i
    C.cell(row=r_, column=2, value=f'=Fact_Product!G{2+i}'); C.cell(row=r_, column=3, value=f'=Fact_Product!H{2+i}'); C.cell(row=r_, column=4, value=f'=Fact_Product!I{2+i}')
    pc = f'{FPr("A")},$B{r_},{FPr("B")},$C$4,{FPr("C")},">="&$C$2,{FPr("C")},"<="&$C$3'
    C.cell(row=r_, column=5, value=f'=SUMIFS({FPr("D")},{pc})').number_format = '#,##0'
    C.cell(row=r_, column=6, value=f'=E{r_}+ROW()*0.000001').number_format = '0.000000'
    C.cell(row=r_, column=7, value=f'=SUMIFS({FPr("E")},{pc})').number_format = '#,##0'
R_P1, R_P2 = R_P + 1, R_P + TOP_PRODUCTS
R_PT = R_P2 + 3
for j, t in enumerate(['Rank', 'Index', 'Label', 'Net revenue', 'Units', 'StockCode']): C.cell(row=R_PT, column=2 + j, value=t).font = font(10, True)
for k in range(10):
    r_ = R_PT + 1 + k
    C.cell(row=r_, column=2, value=k + 1)
    C.cell(row=r_, column=3, value=f'=MATCH(LARGE($F${R_P1}:$F${R_P2},B{r_}),$F${R_P1}:$F${R_P2},0)')
    C.cell(row=r_, column=4, value=f'=IF(INDEX($E${R_P1}:$E${R_P2},C{r_})<=0,"-",LEFT(PROPER(INDEX($D${R_P1}:$D${R_P2},C{r_})),30))')
    C.cell(row=r_, column=5, value=f'=MAX(INDEX($E${R_P1}:$E${R_P2},C{r_}),0)').number_format = '£#,##0,"k"'
    C.cell(row=r_, column=6, value=f'=INDEX($G${R_P1}:$G${R_P2},C{r_})').number_format = '#,##0'
    C.cell(row=r_, column=7, value=f'=INDEX($C${R_P1}:$C${R_P2},C{r_})')
R_PT1, R_PT2 = R_PT + 1, R_PT + 10

# Segments
R_S = R_PT2 + 4; sec(R_S - 1, 'RFM segments (all-time snapshot; follows Country filter)')
for j, t in enumerate(['Segment', 'Customers', '% customers', 'Net revenue', '% revenue', 'Avg rev / customer']): C.cell(row=R_S, column=2 + j, value=t).font = font(10, True)
for i in range(5):
    r_ = R_S + 1 + i
    C.cell(row=r_, column=2, value=f'=Lists!I{2+i}')
    C.cell(row=r_, column=3, value=f'=COUNTIFS({Ur("M")},$B{r_},{Ur("B")},$C$4)').number_format = '#,##0'
    C.cell(row=r_, column=4, value=f'=IF(SUM($C${R_S+1}:$C${R_S+5})=0,0,C{r_}/SUM($C${R_S+1}:$C${R_S+5}))').number_format = '0.0%'
    C.cell(row=r_, column=5, value=f'=SUMIFS({Ur("I")},{Ur("M")},$B{r_},{Ur("B")},$C$4)').number_format = '#,##0'
    C.cell(row=r_, column=6, value=f'=IF(SUM($E${R_S+1}:$E${R_S+5})=0,0,E{r_}/SUM($E${R_S+1}:$E${R_S+5}))').number_format = '0.0%'
    C.cell(row=r_, column=7, value=f'=IF(C{r_}=0,0,E{r_}/C{r_})').number_format = '#,##0'
R_S1, R_S2 = R_S + 1, R_S + 5
r_ = R_S2 + 1; C.cell(row=r_, column=2, value='Total').font = font(10, True)
C.cell(row=r_, column=3, value=f'=SUM(C{R_S1}:C{R_S2})').number_format = '#,##0'; C.cell(row=r_, column=4, value=f'=SUM(D{R_S1}:D{R_S2})').number_format = '0.0%'
C.cell(row=r_, column=5, value=f'=SUM(E{R_S1}:E{R_S2})').number_format = '#,##0'; C.cell(row=r_, column=6, value=f'=SUM(F{R_S1}:F{R_S2})').number_format = '0.0%'
C.cell(row=r_, column=7, value=f'=IF(C{r_}=0,0,E{r_}/C{r_})').number_format = '#,##0'
R_ST = r_

# Top customers
R_C = R_ST + 4; sec(R_C - 1, 'Top 10 customers (follows Country + Year filters)')
for j, t in enumerate(['Rank', 'Index', 'Customer ID', 'Country', 'Net revenue', 'Lifetime orders', 'Last order', 'Segment']): C.cell(row=R_C, column=2 + j, value=t).font = font(10, True)
for k in range(10):
    r_ = R_C + 1 + k
    C.cell(row=r_, column=2, value=k + 1)
    C.cell(row=r_, column=3, value=f'=IF(LARGE({Ur("O")},B{r_})<=0,"",MATCH(LARGE({Ur("O")},B{r_}),{Ur("O")},0))')
    for col, src, fmt in ((4, 'A', '0'), (5, 'B', None), (6, 'N', '#,##0'), (7, 'E', '0'), (8, 'D', 'dd-mmm-yy'), (9, 'M', None)):
        x = C.cell(row=r_, column=col, value=f'=IF($C{r_}="","",INDEX({Ur(src)},$C{r_}))')
        if fmt: x.number_format = fmt
R_C1, R_C2 = R_C + 1, R_C + 10
for row in C.iter_rows(min_row=2):
    for x in row:
        if x.font is None or x.font.name != FN: x.font = font(10)
C.freeze_panes = 'A2'

# ----------------------------------------------------------------------------------
# 4. DASHBOARD
# ----------------------------------------------------------------------------------
D = wsD
D.sheet_view.showGridLines = False; D.sheet_view.zoomScale = 90
D.column_dimensions['A'].width = 2; D.column_dimensions['R'].width = 2
for j in range(2, 18): D.column_dimensions[L(j)].width = 11.5
LASTROW = 84
for r_ in range(1, LASTROW + 1):
    for j in range(1, 19): D.cell(row=r_, column=j).fill = fill(LIGHT)
def rh(r_, h): D.row_dimensions[r_].height = h
for r_, h in ((1, 38), (2, 18), (3, 6), (4, 26), (5, 6), (6, 4), (7, 17), (8, 22), (9, 22), (10, 17), (11, 8)): rh(r_, h)

D.merge_cells('B1:Q1'); D['B1'] = 'Online Retail II  |  Sales Performance Dashboard'
D['B1'].font = font(20, True, WHITE); D['B1'].alignment = Alignment(vertical='center', indent=1)
for j in range(2, 18): D.cell(row=1, column=j).fill = fill(NAVY)
D.merge_cells('B2:Q2'); D['B2'] = f"UK online gift retailer  |  1 Dec 2009 - 9 Dec 2011  |  Currency GBP (£)  |  Net revenue = sales minus cancellations  |  Pick filters in the yellow cells"
D['B2'].font = font(9, False, 'D5DAE3'); D['B2'].alignment = Alignment(vertical='center', indent=1)
for j in range(2, 18): D.cell(row=2, column=j).fill = fill(NAVY)

def label(cell, text):
    D[cell] = text; D[cell].font = font(10, True, NAVY); D[cell].alignment = Alignment(horizontal='right', vertical='center')
def inputcell(rng, value):
    first = rng.split(':')[0]
    D.merge_cells(rng)
    for row in D[rng]:
        for x in row: x.fill = fill('FFF2CC'); x.border = Border(top=Side(style='medium', color=AMBER), bottom=Side(style='medium', color=AMBER), left=Side(style='medium', color=AMBER), right=Side(style='medium', color=AMBER))
    D[first] = value; D[first].font = font(11, True, NAVY); D[first].alignment = Alignment(horizontal='center', vertical='center')
label('B4', 'COUNTRY ▸'); inputcell('C4:E4', 'All Countries')
label('G4', 'YEAR ▸');    inputcell('H4:I4', 'All Years')
label('K4', 'TREND METRIC ▸'); inputcell('L4:M4', 'Net Revenue')
D.merge_cells('N4:Q4'); D['N4'] = '◂ Click a yellow cell, then the drop-down arrow. Everything updates instantly.'
D['N4'].font = font(9, False, GREY, True); D['N4'].alignment = Alignment(wrap_text=True, vertical='center')
dv1 = DataValidation(type='list', formula1=f'=Lists!$A$2:$A${1+NCTRY}', allow_blank=False, showDropDown=False)
dv2 = DataValidation(type='list', formula1='=Lists!$C$2:$C$5', allow_blank=False)
dv3 = DataValidation(type='list', formula1='=Lists!$E$2:$E$5', allow_blank=False)
for dv, cell, ttl, msg in ((dv1, 'C4', 'Country', 'Choose a country or All Countries'), (dv2, 'H4', 'Year', 'Choose a calendar year or All Years'), (dv3, 'L4', 'Trend metric', 'Metric shown in the monthly trend chart')):
    dv.error = 'Please pick a value from the list.'; dv.errorTitle = 'Invalid choice'; dv.promptTitle = ttl; dv.prompt = msg; dv.showErrorMessage = True; dv.showInputMessage = True
    D.add_data_validation(dv); dv.add(cell)

# KPI cards
cards = [('NET REVENUE', '=Calc!$C$23', '[>=1000000]£#,##0.00,,"M";[>=1000]£#,##0.0,"K";£#,##0', '=Calc!$C$48', TEAL),
         ('ORDERS', '=Calc!$C$24', '#,##0', '=Calc!$C$49', NAVY),
         ('AVG ORDER VALUE', '=Calc!$C$26', '£#,##0.00', '=Calc!$C$50', PURPLE),
         ('UNITS SOLD', '=Calc!$C$25', '#,##0', '=Calc!$C$51', GREEN),
         ('ACTIVE CUSTOMERS', '=Calc!$C$28', '#,##0', '=Calc!$C$52', AMBER),
         ('RETURN RATE', '=Calc!$C$27', '0.0%', '=Calc!$C$53', CORAL),
         ('REVENUE / CUSTOMER', '=Calc!$C$31', '£#,##0', '=Calc!$C$54', TEAL),
         ('ITEMS / ORDER', '=Calc!$C$32', '0.0', '=Calc!$C$55', NAVY)]
for i, (t, f, nf, sb, colr) in enumerate(cards):
    c1 = 2 + 2 * i; c2 = c1 + 1; a, b = L(c1), L(c2)
    for r_ in (7, 8, 9, 10):
        D.merge_cells(f'{a}{r_}:{b}{r_}') if r_ in (7, 10) else None
        for cc in (c1, c2):
            x = D.cell(row=r_, column=cc); x.fill = fill(WHITE)
            x.border = Border(top=Side(style='thick', color=colr) if r_ == 7 else None, left=thin if cc == c1 else None, right=thin if cc == c2 else None, bottom=thin if r_ == 10 else None)
    D.merge_cells(f'{a}8:{b}9')
    D[f'{a}7'] = t; D[f'{a}7'].font = font(8, True, GREY); D[f'{a}7'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    D[f'{a}8'] = f; D[f'{a}8'].number_format = nf; D[f'{a}8'].font = font(19, True, NAVY); D[f'{a}8'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
    D[f'{a}10'] = sb; D[f'{a}10'].font = font(8, False, colr if False else GREY); D[f'{a}10'].alignment = Alignment(horizontal='left', vertical='center', indent=1, shrink_to_fit=True)

def caption(rng, formula):
    first = rng.split(':')[0]; D.merge_cells(rng); D[first] = formula
    D[first].font = font(11, True, NAVY); D[first].alignment = Alignment(vertical='center')
    r_ = int(''.join(ch for ch in first if ch.isdigit()))
    for row in D[rng]:
        for x in row: x.border = Border(bottom=Side(style='medium', color=TEAL))
    rh(r_, 20)
caption('B12:I12', '="Monthly trend  |  "&Dashboard!$L$4&"  |  "&Calc!$C$8&"  (all months)"')
caption('J12:Q12', '="Net revenue by calendar month, 2010 vs 2011  |  "&Calc!$C$8&"  (Dec 2011 partial)"')
caption('B30:I30', f'="Top 10 international markets  |  "&Calc!$C$7&"  |  UK = "&TEXT({UKSHARE},"0%")&" of net revenue"')
caption('J30:Q30', '="Top 10 products by net revenue  |  "&Calc!$C$7&"  |  "&Calc!$C$8')
caption('B50:G50', '="Customer segments (RFM, all-time)  |  "&Calc!$C$8')
caption('H50:Q50', '="Top 10 customers  |  "&Calc!$C$7&"  |  "&Calc!$C$8')
caption('B71:Q71', '="When customers buy: gross sales by weekday and hour (£k)  |  "&Calc!$C$7&"  |  "&Calc!$C$8')
for r_ in range(13, 29): rh(r_, 15)
for r_ in range(31, 49): rh(r_, 15)
for r_ in range(51, 70): rh(r_, 15)
for r_ in (6, 29, 49, 70): rh(r_, 8 if r_ != 6 else 4)

# ---- charts
def style_chart(ch, w, h, legend=False):
    ch.width, ch.height = w, h
    if not legend: ch.legend = None
    else: ch.legend.position = 'b'
    ch.graphical_properties = GraphicalProperties(solidFill=WHITE, ln=LineProperties(solidFill='D5DAE3'))
    ch.title = None
def axis_fmt(ax, nf=None, grid=True, size=900):
    ax.delete = False
    if nf: ax.number_format = nf; ax.numFmt.sourceLinked = False
    ax.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(sz=size, solidFill='4A5568', latin=DFont(typeface=FN))), endParaRPr=CharacterProperties())])
    if not grid: ax.majorGridlines = None
    elif ax.majorGridlines is not None: ax.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill='E5E9F0'))

# 1 monthly trend (line)
ch1 = LineChart(); style_chart(ch1, 18.0, 8.4)
ch1.add_data(Reference(C, min_col=10, min_row=R_M, max_row=R_M2), titles_from_data=True)
ch1.set_categories(Reference(C, min_col=9, min_row=R_M1, max_row=R_M2))
s = ch1.series[0]; s.graphicalProperties.line.solidFill = TEAL; s.graphicalProperties.line.width = 28000; s.smooth = False
s.marker.symbol = 'circle'; s.marker.size = 6; s.marker.graphicalProperties = GraphicalProperties(solidFill=TEAL, ln=LineProperties(solidFill=WHITE))
axis_fmt(ch1.x_axis, grid=False, size=800); axis_fmt(ch1.y_axis, '#,##0')
ch1.x_axis.tickLblSkip = 2
D.add_chart(ch1, 'B13')
# 2 YoY (clustered columns)
ch2 = BarChart(); ch2.type = 'col'; ch2.grouping = 'clustered'; style_chart(ch2, 18.0, 8.4, legend=True)
ch2.add_data(Reference(C, min_col=4, max_col=5, min_row=R_Y, max_row=R_Y2), titles_from_data=True)
ch2.set_categories(Reference(C, min_col=3, min_row=R_Y1, max_row=R_Y2))
for s, colr in zip(ch2.series, (GREY, TEAL)): s.graphicalProperties.solidFill = colr; s.graphicalProperties.line.noFill = True
ch2.gapWidth = 60; ch2.overlap = -10
axis_fmt(ch2.x_axis, grid=False); axis_fmt(ch2.y_axis, '£#,##0,"k"')
D.add_chart(ch2, 'J13')
# 3 intl markets (horizontal bar)
def hbar(cats, vals, colr, anchor, w, h):
    ch = BarChart(); ch.type = 'bar'; style_chart(ch, w, h)
    ch.add_data(vals, titles_from_data=False); ch.set_categories(cats)
    s = ch.series[0]; s.graphicalProperties.solidFill = colr; s.graphicalProperties.line.noFill = True
    ch.gapWidth = 45
    ch.dataLabels = DataLabelList(); ch.dataLabels.showVal = True
    for a in ('showSerName', 'showCatName', 'showLegendKey', 'showPercent'): setattr(ch.dataLabels, a, False)
    ch.dataLabels.numFmt = '£#,##0,"k"'; ch.dataLabels.position = 'outEnd'
    ch.x_axis.scaling.orientation = 'maxMin'
    axis_fmt(ch.x_axis, grid=False); axis_fmt(ch.y_axis, '£#,##0,"k"')
    ch.y_axis.delete = True
    D.add_chart(ch, anchor); return ch
hbar(Reference(C, min_col=4, min_row=R_KT1, max_row=R_KT2), Reference(C, min_col=5, min_row=R_KT1, max_row=R_KT2), PURPLE, 'B31', 18.0, 9.2)
hbar(Reference(C, min_col=4, min_row=R_PT1, max_row=R_PT2), Reference(C, min_col=5, min_row=R_PT1, max_row=R_PT2), GREEN, 'J31', 18.0, 9.2)
# 5 segments
ch5 = BarChart(); ch5.type = 'col'; style_chart(ch5, 13.2, 9.7)
ch5.add_data(Reference(C, min_col=3, min_row=R_S1, max_row=R_S2), titles_from_data=False)
ch5.set_categories(Reference(C, min_col=2, min_row=R_S1, max_row=R_S2))
s = ch5.series[0]; s.graphicalProperties.line.noFill = True
for i, colr in enumerate((TEAL, GREEN, PURPLE, AMBER, GREY)):
    pt = DataPoint(idx=i); pt.graphicalProperties = GraphicalProperties(solidFill=colr); s.dPt.append(pt)
ch5.gapWidth = 50
ch5.dataLabels = DataLabelList(); ch5.dataLabels.showVal = True
for a in ('showSerName', 'showCatName', 'showLegendKey', 'showPercent'): setattr(ch5.dataLabels, a, False)
ch5.dataLabels.numFmt = '#,##0'; ch5.dataLabels.position = 'outEnd'
axis_fmt(ch5.x_axis, grid=False, size=800); axis_fmt(ch5.y_axis, '#,##0')
D.add_chart(ch5, 'B51')

# ---- Top customers table (H51:Q61)
hdr = [('H', 'H', '#'), ('I', 'I', 'Customer'), ('J', 'K', 'Country'), ('L', 'M', 'Net revenue'), ('N', 'N', 'Orders'), ('O', 'O', 'Last order'), ('P', 'Q', 'Segment')]
for a, b, t in hdr:
    if a != b: D.merge_cells(f'{a}51:{b}51')
    D[f'{a}51'] = t; D[f'{a}51'].font = font(9, True, WHITE); D[f'{a}51'].alignment = Alignment(horizontal='center', vertical='center')
    for cc in range(ord(a) - 64, ord(b) - 63): D.cell(row=51, column=cc).fill = fill(NAVY)
caption_row = 50
for k in range(10):
    r_ = 52 + k; src = R_C1 + k
    spec = [('H', 'H', f'=Calc!$B${src}', '0'), ('I', 'I', f'=Calc!$D${src}', '0'), ('J', 'K', f'=Calc!$E${src}', None), ('L', 'M', f'=Calc!$F${src}', '£#,##0'),
            ('N', 'N', f'=Calc!$G${src}', '0'), ('O', 'O', f'=Calc!$H${src}', 'dd-mmm-yy'), ('P', 'Q', f'=Calc!$I${src}', None)]
    for a, b, f, nf in spec:
        if a != b: D.merge_cells(f'{a}{r_}:{b}{r_}')
        D[f'{a}{r_}'] = f; D[f'{a}{r_}'].font = font(9, a == 'I'); D[f'{a}{r_}'].alignment = Alignment(horizontal='center', vertical='center')
        if nf: D[f'{a}{r_}'].number_format = nf
        for cc in range(ord(a) - 64, ord(b) - 63): D.cell(row=r_, column=cc).fill = fill(WHITE); D.cell(row=r_, column=cc).border = Border(bottom=thin)
# ---- Segment summary table (H63:Q69)
shdr = [('H', 'J', 'Segment'), ('K', 'K', 'Customers'), ('L', 'L', '% cust.'), ('M', 'N', 'Net revenue'), ('O', 'O', '% rev.'), ('P', 'Q', 'Avg rev / cust.')]
for a, b, t in shdr:
    if a != b: D.merge_cells(f'{a}63:{b}63')
    D[f'{a}63'] = t; D[f'{a}63'].font = font(9, True, WHITE); D[f'{a}63'].alignment = Alignment(horizontal='center', vertical='center')
    for cc in range(ord(a) - 64, ord(b) - 63): D.cell(row=63, column=cc).fill = fill(NAVY)
for k in range(6):
    r_ = 64 + k; src = R_S1 + k; bold = (k == 5)
    spec = [('H', 'J', f'=Calc!$B${src}', None), ('K', 'K', f'=Calc!$C${src}', '#,##0'), ('L', 'L', f'=Calc!$D${src}', '0.0%'),
            ('M', 'N', f'=Calc!$E${src}', '£#,##0'), ('O', 'O', f'=Calc!$F${src}', '0.0%'), ('P', 'Q', f'=Calc!$G${src}', '£#,##0')]
    for a, b, f, nf in spec:
        if a != b: D.merge_cells(f'{a}{r_}:{b}{r_}')
        D[f'{a}{r_}'] = f; D[f'{a}{r_}'].font = font(9, bold); D[f'{a}{r_}'].alignment = Alignment(horizontal='left' if a == 'H' else 'center', vertical='center', indent=1 if a == 'H' else 0)
        if nf: D[f'{a}{r_}'].number_format = nf
        for cc in range(ord(a) - 64, ord(b) - 63): D.cell(row=r_, column=cc).fill = fill(WHITE); D.cell(row=r_, column=cc).border = Border(bottom=thin)
D['H70'] = 'Segments: Champions = bought in the last ~month & 4+ orders | Loyal | Recent / New = recent, 1-3 orders | At Risk = lapsed, 4+ orders | Hibernating / Lost'
D['H70'].font = font(7, False, GREY, True); rh(70, 11)

# ---- Heatmap (row 72 header, 73-79 days)
D['B72'] = 'Day \\ Hour'; D['B72'].font = font(9, True, WHITE); D['B72'].fill = fill(NAVY); D['B72'].alignment = Alignment(horizontal='center')
for j, hr in enumerate(range(6, 21)):
    x = D.cell(row=72, column=3 + j, value=hr); x.number_format = '00":00"'; x.font = font(9, True, WHITE); x.fill = fill(NAVY); x.alignment = Alignment(horizontal='center')
for i, dname in enumerate(WD):
    r_ = 73 + i
    D.cell(row=r_, column=2, value=dname).font = font(10, True, NAVY); D.cell(row=r_, column=2).fill = fill(WHITE); D.cell(row=r_, column=2).alignment = Alignment(horizontal='center')
    for j in range(15):
        cl = L(3 + j)
        x = D.cell(row=r_, column=3 + j, value=f'=SUMIFS({FTr("E")},{FTr("B")},Calc!$C$4,{FTr("A")},">="&Calc!$C$2,{FTr("A")},"<="&Calc!$C$3,{FTr("C")},{i+1},{FTr("D")},{cl}$72)')
        x.number_format = '[=0]"";[<10000]0.0,;#,##0,'; x.font = font(9); x.alignment = Alignment(horizontal='center', vertical='center')
        x.border = Border(left=Side(style='thin', color=WHITE), right=Side(style='thin', color=WHITE), top=Side(style='thin', color=WHITE), bottom=Side(style='thin', color=WHITE))
    rh(r_, 20)
D.conditional_formatting.add('C73:Q79', ColorScaleRule(start_type='min', start_color='FFFFFF', mid_type='percentile', mid_value=60, mid_color='9ED9DE', end_type='max', end_color='0A6F7A'))
D['B80'] = 'Gross sales only (cancellations excluded). Values in £ thousands; blank = no sales. Saturday trading is negligible.'
D['B80'].font = font(8, False, GREY, True)
D['B82'] = 'Data model: Fact_Monthly / Fact_Product / Fact_Time / Customers sheets  |  Logic: Calc sheet  |  Definitions & cleaning log: Notes sheet'
D['B82'].font = font(8, False, GREY)
D.freeze_panes = 'A6'
D.print_area = f'A1:R{LASTROW}'; D.page_setup.orientation = 'landscape'; D.page_setup.fitToWidth = 1; D.page_setup.fitToHeight = 0
D.sheet_properties.pageSetUpPr.fitToPage = True

# ----------------------------------------------------------------------------------
# 5. NOTES
# ----------------------------------------------------------------------------------
N = wsN; N.sheet_view.showGridLines = False
N.column_dimensions['A'].width = 2; N.column_dimensions['B'].width = 34; N.column_dimensions['C'].width = 110
def nrow(r_, a, b=None, head=False):
    x = N.cell(row=r_, column=2, value=a); x.font = font(11 if head else 10, True, WHITE if head else NAVY); x.alignment = Alignment(vertical='top', wrap_text=True)
    if head:
        for cc in (2, 3): N.cell(row=r_, column=cc).fill = fill(NAVY)
    if b is not None:
        y = N.cell(row=r_, column=3, value=b); y.font = font(10); y.alignment = Alignment(vertical='top', wrap_text=True)
    return r_ + 1
r_ = 1
N['B1'] = 'Notes: how to use, definitions, data cleaning and assumptions'; N['B1'].font = font(16, True, NAVY); r_ = 3
r_ = nrow(r_, 'HOW TO USE', head=True)
r_ = nrow(r_, 'Filters', 'On the Dashboard sheet change the three YELLOW cells (Country, Year, Trend metric) using their drop-down arrows. All KPIs, charts and tables recalculate instantly. No macros are needed.')
r_ = nrow(r_, 'What each filter drives', 'Country: every panel except the "Top 10 international markets" chart (which compares countries). Year: KPIs, markets, products, top customers, heat map. The monthly trend and 2010-vs-2011 charts always show all months. Segments follow Country only (they are an all-time snapshot).')
r_ = nrow(r_, 'Sheets', 'Dashboard (the report) | Calc (all formulas feeding the dashboard) | Customers (one row per customer incl. RFM scoring formulas) | Fact_Monthly, Fact_Product, Fact_Time (pre-aggregated data) | Lists (drop-down sources).')
r_ += 1; r_ = nrow(r_, 'DEFINITIONS', head=True)
r_ = nrow(r_, 'Gross sales', 'Quantity x Price on normal (non-cancelled) invoices.')
r_ = nrow(r_, 'Returns / cancellations', 'Absolute value of lines on invoices starting with "C" (customer cancellations / credit notes).')
r_ = nrow(r_, 'Net revenue', 'Gross sales minus returns. Headline revenue figure.')
r_ = nrow(r_, 'Orders', 'Unique non-cancelled invoice numbers.')
r_ = nrow(r_, 'Avg order value', 'Gross sales / orders.')
r_ = nrow(r_, 'Return rate', 'Returns / gross sales (value-based). Returns are dated when processed, so monthly rates can be noisy.')
r_ = nrow(r_, 'Active customers', 'Customers with a Customer ID and positive net revenue in the selected country/year.')
r_ = nrow(r_, 'RFM segments', 'Recency = days from the snapshot date (Calc!C5) to last order, scored 4 (best) to 1 by quartile. Frequency = order count scored 1-4 using editable cut-offs in Calc!C16:C18. Champions: R=4 & F>=3 | Loyal: R=3 & F>=3 | Recent / New: R>=3 & F<=2 | At Risk: R<=2 & F>=3 | Hibernating / Lost: R<=2 & F<=2.')
r_ = nrow(r_, 'Customer country', 'Each customer is assigned the country where they generated the most revenue.')
r_ += 1; r_ = nrow(r_, 'DATA CLEANING LOG (done when the source file was processed)', head=True)
r_ = nrow(r_, 'Rows in source', f"{log['raw_rows']:,} rows across the two year sheets (the stray 'Sheet1' pivot in the source file was ignored).")
r_ = nrow(r_, 'Overlap removed', f"{log['overlap']:,} duplicate rows: 1-9 Dec 2010 appears in both year sheets, so it was kept once.")
r_ = nrow(r_, 'Non-merchandise removed', f"{log['nonmerch']:,} rows: postage, manual adjustments, bank charges, Amazon fees, samples, gift vouchers, test codes, bad-debt adjustment invoices ('A').")
r_ = nrow(r_, 'Zero / negative price removed', f"{log['price0']:,} rows (stock write-offs, damaged goods, free items).")
r_ = nrow(r_, 'Other', f"{log['negqty']:,} non-cancel rows with negative quantity removed.")
r_ = nrow(r_, 'Clean rows analysed', f"{log['clean']:,} rows, {log['dmin']:%d %b %Y} to {log['dmax']:%d %b %Y}.")
r_ += 1; r_ = nrow(r_, 'ASSUMPTIONS & LIMITATIONS', head=True)
r_ = nrow(r_, 'Partial months', 'Dec 2009 contains data from 1 Dec only and Dec 2011 ends on 9 Dec, so Year = 2009 and the last point of 2011 are partial periods.')
r_ = nrow(r_, 'YoY comparison', 'Card sub-texts show YoY only for Year = 2011, comparing Jan to the last complete month (Calc!C6 = 11, i.e. Nov) of 2011 against 2010 to avoid comparing a partial December.')
r_ = nrow(r_, 'Missing Customer ID', f"{log['no_cust_rows']:,} cleaned rows ({log['no_cust_rev_share']:.0%} of net revenue) have no Customer ID. They are included in revenue, orders, units, products and the heat map, but excluded from customer counts, RFM and top customers.")
r_ = nrow(r_, 'Product universe', f"Product rankings cover the top {TOP_PRODUCTS} products by lifetime net revenue ({log['prod_cov']:.0%} of net revenue) to keep the workbook fast.")
r_ = nrow(r_, 'UK dominance', 'The UK is the vast majority of revenue, so the country chart ranks international markets (UK excluded); the UK share is shown in its caption.')
r_ = nrow(r_, 'Refreshing', 'To rebuild from a newer export, run build_dashboard.py (included) against the new source file.')
for rr in range(1, r_ + 1): N.row_dimensions[rr].height = None

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved', OUT)
